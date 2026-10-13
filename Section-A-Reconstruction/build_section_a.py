import json, re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
out=root/'ABLLS_R_Section_A_Reconstructed.docx'
data=json.loads((root/'ocr'/'cell_ocr.json').read_text(encoding='utf-8'))
by={(x['task'],x['field']):x['text'] for x in data}
fields=['task','score','task_name','task_objective','question','examples','criteria','notes']
headers=['TASK','SCORE','TASK NAME','TASK OBJECTIVE','QUESTION','EXAMPLES','CRITERIA','NOTES']
pages={13:range(1,6),14:range(6,10),15:range(10,15),16:range(15,20)}

def clean(task,field,text):
    if field=='task': return task.replace('A','A ')
    if field=='score':
        levels='0 1 2 3 4' if any(ch in text for ch in '34') else '0   1   2'
        return '\n'.join([levels]*4)
    text=re.sub(r'\s+',' ',text).strip()
    fixes={
      ('A10','criteria'):('task: related','task-related'),
      ('A1','task_objective'):('take/ use','take/ use'),
    }
    if (task,field) in fixes: text=text.replace(*fixes[(task,field)])
    return text

doc=Document()
sec=doc.sections[0]
def configure_section(sec):
    sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11); sec.page_height=Inches(8.5)
    sec.left_margin=sec.right_margin=Inches(.32); sec.top_margin=Inches(.30); sec.bottom_margin=Inches(.38)
    sec.header_distance=Inches(.12); sec.footer_distance=Inches(.15)
configure_section(sec)
styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(6.4)
widths=[.39,.56,.89,1.59,1.35,1.31,1.80,2.47]

def borders(table):
    tblPr=table._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'000000'); b.append(e)
    tblPr.append(b)

def margins(cell,top=20,start=28,bottom=20,end=28):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=OxmlElement('w:'+name); node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa'); tcMar.append(node)

def no_cell_borders(cell):
    tcPr=cell._tc.get_or_add_tcPr(); tcBorders=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'nil'); tcBorders.append(e)
    tcPr.append(tcBorders)

def add_text(cell,text,bold_word=None):
    cell.text=''
    parts=text.split('\n') if '\n' in text else [text]
    for i,part in enumerate(parts):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
        if bold_word and bold_word in part:
            a,b=part.split(bold_word,1); p.add_run(a); rr=p.add_run(bold_word); rr.bold=True; p.add_run(b)
        else: p.add_run(part)
        for r in p.runs: r.font.name='Arial'; r.font.size=Pt(6.4); r.font.color.rgb=RGBColor(0,0,0)

for pidx,(pdf_page,tasks) in enumerate(pages.items()):
    if pidx:
        pb=doc.add_paragraph(); pb.paragraph_format.space_after=Pt(0); pb.add_run().add_break(WD_BREAK.PAGE)
    table=doc.add_table(rows=3,cols=8); table.autofit=False; borders(table)
    table.allow_autofit=False
    title_cell=table.cell(0,0).merge(table.cell(0,7)); no_cell_borders(title_cell); margins(title_cell,10,0,40,0)
    add_text(title_cell,'Assessment of Basic Language and Learning Skills - Revised'); title_cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    title_cell.paragraphs[0].runs[0].bold=True; title_cell.paragraphs[0].runs[0].font.size=Pt(11)
    sub_cell=table.cell(1,0).merge(table.cell(1,7)); no_cell_borders(sub_cell); margins(sub_cell,0,0,50,0)
    add_text(sub_cell,'COOPERATION AND REINFORCER EFFECTIVENESS'+(' (Continued)' if pidx else ''))
    sub_cell.paragraphs[0].runs[0].bold=True; sub_cell.paragraphs[0].runs[0].font.size=Pt(8.5)
    for j,h in enumerate(headers):
        c=table.cell(2,j); c.width=Inches(widths[j]); margins(c,10,18,10,18); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(c,h); c.paragraphs[0].runs[0].bold=True; c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    trPr=table.rows[2]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for n in tasks:
        task=f'A{n}'; cells=table.add_row().cells
        for j,field in enumerate(fields):
            c=cells[j]; c.width=Inches(widths[j]); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            word=None
            if task=='A13' and field=='task_name': word='tangible'
            elif task=='A14' and field=='task_name': word='interaction'
            elif task=='A16' and field=='task_name': word='social'
            add_text(c,clean(task,field,by.get((task,field),'')),word)
            if field in ('task','score'): c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        table.rows[-1].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
for sec in doc.sections:
    sec.footer.is_linked_to_previous=False
    footer=sec.footer.paragraphs[0]; footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('Copyright © 1994-2010 By Behavior Analysts, Inc.                                      Page ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
    r=footer.add_run('                                      The ABLLS-R® Protocol')
    for rr in footer.runs: rr.font.name='Arial'; rr.font.size=Pt(6.5)
doc.save(out)
print(out)
