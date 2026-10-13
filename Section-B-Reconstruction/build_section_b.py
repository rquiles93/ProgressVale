import json, re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
out = root / 'ABLLS_R_Section_B_Reconstructed.docx'
data = json.loads((root / 'ocr' / 'cell_ocr.json').read_text(encoding='utf-8'))
by = {(x['task'], x['field']): x['text'] for x in data}
fields = ['task','score','task_name','task_objective','question','examples','criteria','notes']
headers = ['TASK','SCORE','TASK NAME','TASK OBJECTIVE','QUESTION','EXAMPLES','CRITERIA','NOTES']
pages = {17: range(1,6), 18: range(6,12), 19: range(12,17),
         20: range(17,21), 21: range(21,25), 22: range(25,28)}
graphic_examples = {'B1':17, 'B10':18, 'B11':18, 'B14':19, 'B15':19, 'B25':22}

def clean(task, field, text):
    if field == 'task':
        return task.replace('B', 'B ')
    if field == 'score':
        levels = '0 1 2 3 4' if any(ch in text for ch in '34') else '0   1   2'
        return '\n'.join([levels] * 4)
    text = re.sub(r'\s+', ' ', text).strip()
    replacements = {
        ('B17','question'): [('e.qg.,', 'e.g.,')],
        ('B22','examples'): [('ared', 'a red')],
        ('B25','question'): [("1283's", "123's")],
        ('B27','criteria'): [('candraw', 'can draw'), ('line ina', 'line in a')],
    }
    for old, new in replacements.get((task, field), []):
        text = text.replace(old, new)
    return text

def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node = OxmlElement('w:' + edge)
        node.set(qn('w:val'), 'single'); node.set(qn('w:sz'), '5'); node.set(qn('w:color'), '000000')
        borders.append(node)
    tbl_pr.append(borders)

def set_margins(cell, top=20, start=28, bottom=20, end=28):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar'); tc_pr.append(tc_mar)
    for name, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = OxmlElement('w:' + name); node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)

def no_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node = OxmlElement('w:' + edge); node.set(qn('w:val'), 'nil'); borders.append(node)
    tc_pr.append(borders)

def add_text(cell, text, bold_word=None):
    cell.text = ''
    for i, part in enumerate(text.split('\n')):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1
        if bold_word and bold_word in part:
            before, after = part.split(bold_word, 1)
            p.add_run(before); run = p.add_run(bold_word); run.bold = True; p.add_run(after)
        else:
            p.add_run(part)
        for run in p.runs:
            run.font.name = 'Arial'; run.font.size = Pt(6.4); run.font.color.rgb = RGBColor(0,0,0)

def add_graphic(cell, task, page):
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    path = root / 'cells' / f'{task}_examples_p{page}.png'
    p.add_run().add_picture(str(path), width=Inches(1.05))

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11); section.page_height = Inches(8.5)
section.left_margin = section.right_margin = Inches(.32)
section.top_margin = Inches(.30); section.bottom_margin = Inches(.38)
section.header_distance = Inches(.12); section.footer_distance = Inches(.15)
doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(6.4)
widths = [.39,.56,.89,1.59,1.35,1.31,1.80,2.47]

for page_index, (pdf_page, tasks) in enumerate(pages.items()):
    if page_index:
        paragraph = doc.add_paragraph(); paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        # A minimal anchor paragraph prevents Word from positioning the next
        # table's borderless title rows above the printable area.
        anchor = doc.add_paragraph(); anchor.paragraph_format.space_before = Pt(0); anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(12); anchor.add_run(' ').font.size = Pt(1)
    table = doc.add_table(rows=3, cols=8); table.autofit = False; table.allow_autofit = False; set_borders(table)
    title = table.cell(0,0).merge(table.cell(0,7)); no_borders(title); set_margins(title,10,0,40,0)
    add_text(title, 'Assessment of Basic Language and Learning Skills - Revised')
    title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraphs[0].runs[0].bold = True; title.paragraphs[0].runs[0].font.size = Pt(11)
    subtitle = table.cell(1,0).merge(table.cell(1,7)); no_borders(subtitle); set_margins(subtitle,0,0,50,0)
    add_text(subtitle, 'VISUAL PERFORMANCE' + (' (Continued)' if page_index else ''))
    subtitle.paragraphs[0].runs[0].bold = True; subtitle.paragraphs[0].runs[0].font.size = Pt(8.5)
    for col, heading in enumerate(headers):
        cell = table.cell(2,col); cell.width = Inches(widths[col]); set_margins(cell,10,18,10,18)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; add_text(cell, heading)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; cell.paragraphs[0].runs[0].bold = True
    tr_pr = table.rows[2]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); tr_pr.append(repeat)
    for number in tasks:
        task = f'B{number}'; cells = table.add_row().cells
        for col, field in enumerate(fields):
            cell = cells[col]; cell.width = Inches(widths[col]); set_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if field == 'examples' and task in graphic_examples:
                add_graphic(cell, task, graphic_examples[task])
            else:
                bold_word = 'objects' if task == 'B3' and field == 'task_name' else ('pictures' if task == 'B5' and field == 'task_name' else None)
                add_text(cell, clean(task, field, by.get((task,field), '')), bold_word)
            if field in ('task','score'):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

for section in doc.sections:
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]; footer.clear(); footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Copyright © 1994-2010 By Behavior Analysts, Inc.                                      Page ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.add_run('                                      The ABLLS-R® Protocol')
    for run in footer.runs:
        run.font.name = 'Arial'; run.font.size = Pt(6.5)

doc.save(out)
print(out)
