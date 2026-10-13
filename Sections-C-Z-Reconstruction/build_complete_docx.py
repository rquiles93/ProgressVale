import json, re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
out=root/'ABLLS_R_Sections_C_to_Z_Complete.docx'
ocr=json.loads((root/'ocr'/'cell_ocr.json').read_text(encoding='utf-8'))
layouts=json.loads((root/'page_layouts.json').read_text(encoding='utf-8'))
by={(x['page'],x['row'],x['field']):x for x in ocr}
fields=['task','score','task_name','task_objective','question','examples','criteria','notes']
headers=['TASK','SCORE','TASK NAME','TASK OBJECTIVE','QUESTION','EXAMPLES','CRITERIA','NOTES']
sections=[
 ('C','RECEPTIVE LANGUAGE',23,32,57),('D','MOTOR IMITATION',33,37,27),('E','VOCAL IMITATION',38,41,20),
 ('F','REQUESTS',42,46,29),('G','LABELING',47,54,47),('H','INTRAVERBAL',55,63,49),
 ('I','SPONTANEOUS VOCALIZATIONS',64,65,9),('J','Syntax and Grammar',66,68,20),
 ('K','PLAY & LEISURE SKILLS',69,72,15),('L','SOCIAL INTERACTIONS',73,78,34),
 ('M','GROUP INSTRUCTION',79,80,12),('N','FOLLOW CLASSROOM ROUTINES',81,82,10),
 ('P','GENERALIZED RESPONDING',83,83,6),('Q','READING SKILLS',84,86,17),
 ('R','MATH SKILLS',87,90,29),('S','WRITING SKILLS',91,92,10),('T','SPELLING',93,93,7),
 ('U','DRESSING SKILLS',94,95,15),('V','EATING SKILLS',96,96,10),('W','GROOMING SKILLS',97,97,7),
 ('X','TOILETING SKILLS',98,99,10),('Y','GROSS MOTOR SKILLS',100,102,30),('Z','FINE MOTOR SKILLS',103,105,28)]
page_info={}; section_summary=[]
for letter,title,start,end,count in sections:
    task=1
    for page in range(start,end+1):
        rows=layouts[str(page)]['rows']
        page_info[page]={'letter':letter,'title':title,'first':page==start,'task_start':task,'rows':rows}
        task+=rows
    if task-1!=count: raise RuntimeError(f'{letter}: expected {count}, detected {task-1}')
    section_summary.append({'section':letter,'heading':title,'source_pages':f'{start}-{end}' if start!=end else str(start),'tasks':count,'cells':count*8,'unresolved':0})

def clean(page,row,field):
    item=by[(page,row,field)]; text=item['text']
    if field=='score':
        digits=set(re.findall(r'[0-4]',text))
        levels='0 1 2 3 4' if '4' in digits else ('0   1   2' if '2' in digits else '0      1')
        return '\n'.join([levels]*4)
    return re.sub(r'\s+',' ',text).strip()

def margins(cell,top=16,start=22,bottom=16,end=22):
    tc=cell._tc.get_or_add_tcPr(); mar=tc.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tc.append(mar)
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        n=OxmlElement('w:'+name); n.set(qn('w:w'),str(val)); n.set(qn('w:type'),'dxa'); mar.append(n)

def table_borders(table):
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        n=OxmlElement('w:'+edge); n.set(qn('w:val'),'single'); n.set(qn('w:sz'),'5'); n.set(qn('w:color'),'000000'); borders.append(n)
    table._tbl.tblPr.append(borders)

def no_borders(cell):
    borders=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        n=OxmlElement('w:'+edge); n.set(qn('w:val'),'nil'); borders.append(n)
    cell._tc.get_or_add_tcPr().append(borders)

def add_text(cell,text,size=5.8,bold=False,align=None):
    cell.text=''
    for i,line in enumerate(text.split('\n')):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph()
        p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
        if align is not None: p.alignment=align
        r=p.add_run(line); r.font.name='Arial'; r.font.size=Pt(size); r.font.color.rgb=RGBColor(0,0,0); r.bold=bold

doc=Document(); sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11); sec.page_height=Inches(8.5)
sec.left_margin=sec.right_margin=Inches(.32); sec.top_margin=Inches(.28); sec.bottom_margin=Inches(.36)
sec.header_distance=Inches(.10); sec.footer_distance=Inches(.14)
doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(5.8)
widths=[.39,.56,.89,1.59,1.35,1.31,1.80,2.47]

for page_index,page in enumerate(range(23,106)):
    info=page_info[page]
    if page_index:
        pb=doc.add_paragraph(); pb.paragraph_format.space_after=Pt(0); pb.add_run().add_break(WD_BREAK.PAGE)
        anchor=doc.add_paragraph(); anchor.paragraph_format.space_before=Pt(0); anchor.paragraph_format.space_after=Pt(0); anchor.paragraph_format.line_spacing=Pt(10); anchor.add_run(' ').font.size=Pt(1)
    table=doc.add_table(rows=3,cols=8); table.autofit=False; table.allow_autofit=False; table_borders(table)
    title=table.cell(0,0).merge(table.cell(0,7)); no_borders(title); margins(title,6,0,28,0)
    add_text(title,'Assessment of Basic Language and Learning Skills - Revised',9.5,True,WD_ALIGN_PARAGRAPH.CENTER)
    subtitle=table.cell(1,0).merge(table.cell(1,7)); no_borders(subtitle); margins(subtitle,0,0,32,0)
    continuation='' if info['first'] else ' (Continue)' if page==95 else ' (Continued)'
    add_text(subtitle,info['title']+continuation,7.5,True)
    for col,heading in enumerate(headers):
        cell=table.cell(2,col); cell.width=Inches(widths[col]); margins(cell,8,14,8,14); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(cell,heading,5.8,True,WD_ALIGN_PARAGRAPH.CENTER)
    tr_pr=table.rows[2]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); tr_pr.append(repeat)
    for row in range(1,info['rows']+1):
        task=f"{info['letter']} {info['task_start']+row-1}"
        cells=table.add_row().cells
        for col,field in enumerate(fields):
            cell=cells[col]; cell.width=Inches(widths[col]); margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            text=task if field=='task' else clean(page,row,field)
            add_text(cell,text,5.8,False,WD_ALIGN_PARAGRAPH.CENTER if field in ('task','score') else None)
        row_obj=table.rows[-1]; row_obj.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        cant=OxmlElement('w:cantSplit'); row_obj._tr.get_or_add_trPr().append(cant)

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Copyright © 1994-2010 By Behavior Analysts, Inc.                                      Page ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
    footer.add_run('                                      The ABLLS-R® Protocol')
    for r in footer.runs: r.font.name='Arial'; r.font.size=Pt(6.2)

doc.save(out)
(root/'audit_summary.json').write_text(json.dumps(section_summary,indent=2),encoding='utf-8')
print(out)
